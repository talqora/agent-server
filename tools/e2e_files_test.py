"""文件类型端到端:PDF(PyMuPDF)与 DOCX(python-docx)真实解析链路 + 类型白名单拒绝。

覆盖 e2e_test.py(md)之外的文件路径:
  1. 生成真实 PDF(中文,内嵌 CJK 字体)与 DOCX → 上传 → 摄取 SSE 到 run_completed → ready + 分片;
  2. 不支持的类型(.exe)→ 400;
  3. 跨文档对话:检索 citations 指向 PDF/DOCX 之一;
  4. 清理删除。

前置:中间件 + llm-stub + rag-http + rag-worker 已启动。
用法:``uv run python tools/e2e_files_test.py``
"""

from __future__ import annotations

import argparse
import io
import json
import sys
import uuid
from collections.abc import Iterator
from typing import Any

import httpx

PASS: list[str] = []
FAIL: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASS if ok else FAIL).append(name)
    print(f"  {'✓' if ok else '✗'} {name}" + (f"  [{detail}]" if detail and not ok else ""))


def sse_events(response: httpx.Response) -> Iterator[tuple[str | None, dict[str, Any]]]:
    event: str | None = None
    data_lines: list[str] = []
    for line in response.iter_lines():
        if line == "":
            if data_lines:
                raw = "\n".join(data_lines)
                try:
                    payload = json.loads(raw)
                except json.JSONDecodeError:
                    payload = {"raw": raw}
                yield event, payload
            event = None
            data_lines = []
            continue
        if line.startswith(":"):
            continue
        field, _, value = line.partition(":")
        value = value[1:] if value.startswith(" ") else value
        if field == "event":
            event = value
        elif field == "data":
            data_lines.append(value)


def stream_run(client: httpx.Client, run_id: str, headers: dict[str, str]) -> list[str]:
    events: list[str] = []
    with client.stream("GET", f"/runs/{run_id}/stream", headers=headers, timeout=120) as resp:
        for event, _ in sse_events(resp):
            events.append(event or "")
            if event in ("run_completed", "run_failed"):
                break
    return events


def make_pdf() -> bytes:
    import pymupdf

    text = "端到端测试 PDF 内容:这是一份用于验证 PDF 解析链路的文档。" * 40
    doc = pymupdf.open()
    page = doc.new_page()
    try:
        # PyMuPDF 内置简体中文字体(抽取文本可还原)
        page.insert_text((72, 100), text, fontname="china-s", fontsize=10)
    except Exception:  # noqa: BLE001 - 字体不可用时退化为 ASCII(仍验证解析链路)
        page.insert_text((72, 100), "E2E PDF content for parsing. " * 60, fontsize=10)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("端到端测试 DOCX", level=1)
    doc.add_paragraph("这是用于验证 DOCX 解析链路的文档段落。" * 40)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(
    client: httpx.Client,
    headers: dict[str, str],
    filename: str,
    content: bytes,
    mime: str,
) -> httpx.Response:
    return client.post(
        "/documents", headers=headers, files={"file": (filename, content, mime)}
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", default="http://localhost:3101/api")
    args = parser.parse_args()
    client = httpx.Client(base_url=args.base, timeout=120)

    username = f"files_{uuid.uuid4().hex[:8]}"
    token = client.post(
        "/auth/register",
        json={"username": username, "password": "password123", "displayName": "Files"},
    ).json()["token"]
    headers = {"Authorization": f"Bearer {token}"}

    print("== 1. PDF 真实解析链路 ==")
    pdf_up = upload(client, headers, "测试报告.pdf", make_pdf(), "application/pdf")
    check("PDF 上传 201", pdf_up.status_code == 201, pdf_up.text[:200])
    pdf_doc_id = pdf_up.json()["documentId"]
    pdf_events = stream_run(client, pdf_up.json()["runId"], headers)
    check(
        "PDF 摄取完成(parsing→embedding→run_completed)",
        pdf_events[:1] == ["run_started"]
        and "step" in pdf_events
        and pdf_events[-1] == "run_completed",
        str(pdf_events),
    )
    pdf_doc = client.get(f"/documents/{pdf_doc_id}", headers=headers).json()
    check(
        "PDF 文档 ready 且有分片",
        pdf_doc["status"] == "ready" and pdf_doc["chunkCount"] > 0,
        json.dumps(pdf_doc, ensure_ascii=False)[:200],
    )

    print("== 2. DOCX 真实解析链路 ==")
    docx_up = upload(
        client,
        headers,
        "测试文档.docx",
        make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    check("DOCX 上传 201", docx_up.status_code == 201, docx_up.text[:200])
    docx_doc_id = docx_up.json()["documentId"]
    docx_events = stream_run(client, docx_up.json()["runId"], headers)
    check("DOCX 摄取以 run_completed 收尾", docx_events[-1] == "run_completed", str(docx_events[-3:]))
    docx_doc = client.get(f"/documents/{docx_doc_id}", headers=headers).json()
    check(
        "DOCX 文档 ready 且有分片",
        docx_doc["status"] == "ready" and docx_doc["chunkCount"] > 0,
        json.dumps(docx_doc, ensure_ascii=False)[:200],
    )

    print("== 3. 类型白名单拒绝 ==")
    bad = upload(client, headers, "恶意.exe", b"MZ\x90\x00", "application/octet-stream")
    check("不支持类型 400", bad.status_code == 400, str(bad.status_code))

    print("== 4. 跨文档检索对话(引用指向 PDF/DOCX) ==")
    conv = client.post("/conversations", headers=headers, json={}).json()
    done_payload: dict[str, Any] | None = None
    with client.stream(
        "POST",
        f"/conversations/{conv['id']}/messages",
        headers=headers,
        json={"query": "这些文档讲了什么?", "topK": 6},
        timeout=120,
    ) as resp:
        for event, payload in sse_events(resp):
            if event == "done":
                done_payload = payload
                break
            if event == "error":
                check("对话无错误帧", False, json.dumps(payload, ensure_ascii=False))
                break
    citations = (done_payload or {}).get("citations", [])
    cited_ids = {c.get("documentId") for c in citations}
    check(
        "citations 命中 PDF/DOCX",
        bool(cited_ids & {pdf_doc_id, docx_doc_id}),
        f"citations={citations[:3]}",
    )

    print("== 5. 清理 ==")
    for doc_id in (pdf_doc_id, docx_doc_id):
        resp = client.delete(f"/documents/{doc_id}", headers=headers)
        check(f"删文档 {doc_id} 204", resp.status_code == 204, str(resp.status_code))

    print()
    print(f"结果: {len(PASS)} 通过, {len(FAIL)} 失败")
    if FAIL:
        print("失败项:")
        for name in FAIL:
            print(f"  - {name}")
        return 1
    print("文件类型端到端通过 ✅")
    return 0


if __name__ == "__main__":
    sys.exit(main())
